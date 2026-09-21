"""Extract the public Alpha Strike Converter workbook into reviewable JSON.

The workbook is a corroborating/provisional source. This importer preserves its
metadata and does not promote values to canonical catalog data.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--output", default="tools/alpha-strike-workbook-data.json")
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    sheet = workbook["Data"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values[1]:
            continue
        row = {headers[index]: values[index] for index in range(min(len(headers), len(values))) if headers[index]}
        row["source"] = "Alpha Strike Converter v2.3 workbook"
        row["source_creator"] = "Keith Hann"
        row["source_status"] = "provisional"
        rows.append(row)

    output = {
        "source_file": str(Path(args.workbook).name),
        "source_creator": "Keith Hann",
        "source_last_modified_by": "HMS Dreadnought",
        "source_created": "2015-06-09T18:31:49Z",
        "source_modified": "2023-03-07T20:31:51Z",
        "source_status": "provisional",
        "sheet": "Data",
        "row_count": len(rows),
        "rows": rows,
    }
    Path(args.output).write_text(json.dumps(output, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"row_count": len(rows), "output": args.output}, indent=2))


if __name__ == "__main__":
    main()